"""Builders for small CV test files (PDF and DOCX), created in memory so no binaries are committed."""

import io
import zipfile

from docx import Document
from pypdf import PdfWriter

CV_LINES = [
    "Jane Doe",
    "IT Support Specialist",
    "jane.doe@example.com | +254 700 000000 | Nairobi, Kenya",
    "Summary",
    "IT support professional with four years of experience in helpdesk and networking.",
    "Experience",
    "IT Support Officer, Acme Ltd, Nairobi (Mar 2021 - Present)",
    "- Resolved helpdesk tickets and escalated complex faults, handling about 30 tickets per day",
    "- Set up and maintained user accounts in Active Directory",
    "Intern, Beta Networks (2020 - Feb 2021)",
    "- Assisted with network cabling and router configuration (Networking)",
    "Education",
    "BSc Computer Science, University of Nairobi, 2016 - 2020",
    "Certifications",
    "CompTIA A+, CompTIA, 2021",
    "Skills",
    "Helpdesk, Active Directory, Networking, Customer service",
]


def _escape(text: str) -> str:
    return text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def pdf_bytes(lines: list[str] | None = None) -> bytes:
    """A one-page, text-based PDF (hand-built so it needs no extra library)."""
    lines = lines if lines is not None else CV_LINES
    body = (
        "BT /F1 11 Tf 14 TL 50 800 Td " + " ".join(f"({_escape(x)}) Tj T*" for x in lines) + " ET"
    )
    stream = body.encode("latin-1", "replace")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 842] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for number, obj in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n".encode() + obj + b"\nendobj\n"
    xref_at = len(out)
    out += f"xref\n0 {len(objects) + 1}\n".encode() + b"0000000000 65535 f \n"
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_at}\n%%EOF\n".encode()
    return bytes(out)


def blank_pdf_bytes() -> bytes:
    """A valid PDF with a page but no text, like a scan."""
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=842)
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def encrypted_pdf_bytes() -> bytes:
    writer = PdfWriter(clone_from=io.BytesIO(pdf_bytes()))
    writer.encrypt("secret")
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def docx_bytes(lines: list[str] | None = None, *, table: list[list[str]] | None = None) -> bytes:
    document = Document()
    for line in lines if lines is not None else CV_LINES:
        document.add_paragraph(line)
    if table:
        grid = document.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, value in enumerate(row):
                grid.cell(r, c).text = value
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def zip_without_document_xml() -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("hello.txt", "not a docx")
    return buffer.getvalue()


def docx_zip_bomb() -> bytes:
    """Looks like a DOCX but expands to far more than the allowed size."""
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", b"0" * (60 * 1024 * 1024))
    return buffer.getvalue()
