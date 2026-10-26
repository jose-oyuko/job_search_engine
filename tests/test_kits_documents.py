"""Rendering: facts from the database, file names, Word templates, ATS-friendly PDFs (spec §12.6)."""

import io
import re
from typing import Any

import pytest
import time_machine
from docx import Document

from apps.accounts.models import UserSettings
from apps.kits import documents, docx_templates
from apps.kits.models import Kit
from apps.profiles import services as profile_services
from tests.kit_helpers import (
    PDF_AVAILABLE,
    Ids,
    good_cv,
    letter_text,
    make_kit_job,
    make_kit_user,
    requires_pdf,
)

pytestmark = pytest.mark.django_db


@pytest.fixture
def kit_world() -> tuple[Kit, Ids]:
    user, ids = make_kit_user()
    kit = Kit.objects.create(
        user=user,
        job=make_kit_job(),
        status=Kit.Status.READY,
        cv_data=good_cv(ids),
        cover_letter_text=letter_text(),
    )
    return kit, ids


def docx_paragraphs(data: bytes) -> list[str]:
    return [p.text for p in Document(io.BytesIO(data)).paragraphs]


def docx_styles(data: bytes) -> list[tuple[str, str]]:
    return [(p.style.name, p.text) for p in Document(io.BytesIO(data)).paragraphs]


# --- context: facts come from the database ------------------------------------------------------


def test_the_cv_context_is_built_from_profile_rows(kit_world) -> None:  # type: ignore[no-untyped-def]
    kit, _ = kit_world
    ctx = documents.build_cv_context(kit)
    assert ctx["name"] == "Jane Wanjiru Doe"
    assert ctx["contact_line"] == "+254700000000 | jane@example.com | Nairobi, Kenya"
    first = ctx["experiences"][0]
    assert (first["title"], first["company"]) == ("IT Support Technician", "Acme Ltd")
    assert first["dates"] == "Mar 2021 - Present" and first["index"] == 0
    assert ctx["skills"] == ["Helpdesk", "Excel"] and ctx["skills_line"] == "Helpdesk, Excel"
    assert ctx["projects"][0]["name"] == "Inventory tracker"
    assert ctx["education"][0]["title"] == "Diploma in IT"
    assert ctx["education"][0]["institution"] == "Nairobi Polytechnic"
    assert ctx["certifications"][0]["name"] == "CompTIA A+"


def test_links_in_the_profile_appear_in_the_contact_line(kit_world) -> None:  # type: ignore[no-untyped-def]
    kit, _ = kit_world
    profile = profile_services.get_profile(kit.user)
    profile.links = [{"label": "GitHub", "url": "https://github.com/jane"}, "junk"]
    profile.save()
    assert documents.build_cv_context(kit)["contact_line"].endswith("| https://github.com/jane")


def test_ids_from_another_user_never_resolve(kit_world) -> None:  # type: ignore[no-untyped-def]
    kit, _ = kit_world
    _, other = make_kit_user(full_name="Other Person")
    kit.cv_data = good_cv(other)  # another user's IDs, however they got there
    ctx = documents.build_cv_context(kit)
    assert ctx["experiences"] == [] and ctx["skills"] == [] and ctx["education"] == []


def test_a_missing_name_falls_back_to_a_placeholder(kit_world) -> None:  # type: ignore[no-untyped-def]
    kit, _ = kit_world
    kit.user.full_name = ""
    assert documents.build_cv_context(kit)["name"] == "Your name"


def test_dates_cover_current_closed_and_unknown_ranges() -> None:
    import datetime

    d = datetime.date
    assert documents.date_range(d(2021, 3, 1), None, current=True) == "Mar 2021 - Present"
    assert (
        documents.date_range(d(2021, 3, 1), d(2023, 6, 1), current=False) == "Mar 2021 - Jun 2023"
    )
    assert documents.date_range(None, d(2023, 6, 1), current=False) == "Jun 2023"
    assert documents.date_range(None, None, current=False) == ""


def test_the_letter_context_uses_the_users_local_date(kit_world) -> None:  # type: ignore[no-untyped-def]
    kit, _ = kit_world
    UserSettings.objects.filter(user=kit.user).update(timezone="Africa/Nairobi")
    with time_machine.travel("2026-03-04 22:30:00+00:00", tick=False):
        ctx = documents.build_letter_context(kit)
    assert ctx["date"] == "5 March 2026"  # already the 5th in Nairobi (UTC+3)
    assert ctx["company"] == "Globex Corp" and len(ctx["paragraphs"]) == 3
    assert ctx["salutation"] == "Dear Hiring Team," and ctx["closing"] == "Kind regards,"


def test_an_unknown_timezone_does_not_stop_the_letter(kit_world) -> None:  # type: ignore[no-untyped-def]
    kit, _ = kit_world
    UserSettings.objects.filter(user=kit.user).update(timezone="Not/AZone")
    assert documents.build_letter_context(kit)["date"]


def test_paragraph_splitting() -> None:
    assert documents.split_paragraphs("One.\n\n\nTwo.\n  \nThree.") == ["One.", "Two.", "Three."]
    assert documents.split_paragraphs("   ") == []


# --- file names ------------------------------------------------------------------------------------


@pytest.mark.parametrize(
    ("name", "company", "key", "expected"),
    [
        ("Jane Wanjiru Doe", "Globex Corp", "cv_pdf", "Jane_Doe_CV_Globex-Corp.pdf"),
        ("Jane Wanjiru Doe", "Globex Corp", "cv_docx", "Jane_Doe_CV_Globex-Corp.docx"),
        ("Jane Doe", "Globex", "cover_letter_pdf", "Jane_Doe_CoverLetter_Globex.pdf"),
        ("Zoë Müller", "Société Générale", "cv_pdf", "Zoe_Muller_CV_Societe-Generale.pdf"),
        ("Cher", "Acme", "cv_pdf", "Cher_CV_Acme.pdf"),
        ("", "Acme", "cv_pdf", "Candidate_CV_Acme.pdf"),
        ("Jane Doe", "../../etc/passwd", "cv_pdf", "Jane_Doe_CV_etc-passwd.pdf"),
        ("Jane Doe", "日本語", "cv_pdf", "Jane_Doe_CV.pdf"),
    ],
)
def test_download_file_names(name: str, company: str, key: str, expected: str) -> None:
    user, _ = make_kit_user(full_name=name)
    kit = Kit.objects.create(user=user, job=make_kit_job(company=company))
    result = documents.download_filename(kit, key)
    assert result == expected
    assert result.isascii() and re.fullmatch(r"[A-Za-z0-9_.-]+", result)


def test_long_company_names_are_cut() -> None:
    user, _ = make_kit_user()
    kit = Kit.objects.create(user=user, job=make_kit_job(company="A" * 200))
    assert len(documents.download_filename(kit, "cv_pdf")) < 60


# --- HTML (the PDF source) --------------------------------------------------------------------------


def test_the_cv_html_is_single_column_and_plain(kit_world) -> None:  # type: ignore[no-untyped-def]
    kit, _ = kit_world
    html = documents.render_html(documents.DOC_CV, documents.build_cv_context(kit))
    for forbidden in ("<table", "<img", "<svg", "<canvas", "float:", "columns:", "display: flex"):
        assert forbidden not in html
    assert "doc.css" in html
    headings = re.findall(r"<h2>(.*?)</h2>", html)
    assert headings == [
        "Summary",
        "Skills",
        "Experience",
        "Projects",
        "Education",
        "Certifications",
    ]
    css = (documents.TEMPLATE_DIR / "doc.css").read_text(encoding="utf-8")
    assert "float:" not in css and "column" not in css and "flex" not in css and "grid" not in css
    assert "size: A4" in css


def test_generated_text_is_escaped_in_html(kit_world) -> None:  # type: ignore[no-untyped-def]
    kit, ids = kit_world
    cv = good_cv(ids)
    cv["experiences"][0]["bullets"][0]["text"] = '<script>alert("x")</script> & more'
    kit.cv_data = cv
    kit.cover_letter_text = "<b>bold</b> letter"
    cv_html = documents.render_html(documents.DOC_CV, documents.build_cv_context(kit))
    letter_html = documents.render_html(documents.DOC_COVER, documents.build_letter_context(kit))
    assert "<script>" not in cv_html and "&lt;script&gt;" in cv_html
    assert "<b>bold</b>" not in letter_html and "&lt;b&gt;bold&lt;/b&gt;" in letter_html


def test_empty_sections_are_left_out(kit_world) -> None:  # type: ignore[no-untyped-def]
    kit, ids = kit_world
    kit.cv_data = good_cv(
        ids, projects=[], education_ids=[], certification_ids=[], skills=[], summary=""
    )
    html = documents.render_html(documents.DOC_CV, documents.build_cv_context(kit))
    assert re.findall(r"<h2>(.*?)</h2>", html) == ["Experience"]


# --- Word documents ---------------------------------------------------------------------------------


def test_the_cv_docx_has_the_standard_headings_in_order(kit_world) -> None:  # type: ignore[no-untyped-def]
    kit, _ = kit_world
    data = documents.render_docx(documents.DOC_CV, documents.build_cv_context(kit))
    styles = docx_styles(data)
    headings = [text for style_name, text in styles if style_name == "Heading 1"]
    assert headings == [
        "Summary",
        "Skills",
        "Experience",
        "Projects",
        "Education",
        "Certifications",
    ]
    text = "\n".join(t for _, t in styles)
    assert "Jane Wanjiru Doe" in text and "IT Support Technician, Acme Ltd" in text
    assert "Resolved support tickets for 120 staff and cut the backlog by 40%." in text
    bullets = [t for s, t in styles if s == "List Bullet"]
    assert "Set up Windows laptops for new staff." in bullets


def test_no_template_tags_survive_rendering(kit_world) -> None:  # type: ignore[no-untyped-def]
    kit, _ = kit_world
    for doc, ctx in (
        (documents.DOC_CV, documents.build_cv_context(kit)),
        (documents.DOC_COVER, documents.build_letter_context(kit)),
    ):
        text = "\n".join(docx_paragraphs(documents.render_docx(doc, ctx)))
        assert "{{" not in text and "{%" not in text and "}}" not in text


def test_empty_sections_vanish_from_the_docx(kit_world) -> None:  # type: ignore[no-untyped-def]
    kit, ids = kit_world
    kit.cv_data = good_cv(ids, projects=[], education_ids=[], certification_ids=[], skills=[])
    data = documents.render_docx(documents.DOC_CV, documents.build_cv_context(kit))
    headings = [t for s, t in docx_styles(data) if s == "Heading 1"]
    assert headings == ["Summary", "Experience"]


def test_docx_text_is_escaped(kit_world) -> None:  # type: ignore[no-untyped-def]
    kit, ids = kit_world
    cv = good_cv(ids)
    cv["experiences"][0]["bullets"][0]["text"] = 'R&D <fixes> for "staff"'
    kit.cv_data = cv
    text = "\n".join(
        docx_paragraphs(documents.render_docx(documents.DOC_CV, documents.build_cv_context(kit)))
    )
    assert 'R&D <fixes> for "staff"' in text


def test_the_cover_letter_docx_has_greeting_body_and_sign_off(kit_world) -> None:  # type: ignore[no-untyped-def]
    kit, _ = kit_world
    paragraphs = docx_paragraphs(
        documents.render_docx(documents.DOC_COVER, documents.build_letter_context(kit))
    )
    assert paragraphs[0] == "Jane Wanjiru Doe"
    assert "Dear Hiring Team," in paragraphs and "Kind regards," in paragraphs
    assert paragraphs[-1] == "Jane Wanjiru Doe"


def test_the_committed_templates_match_the_builder(tmp_path) -> None:  # type: ignore[no-untyped-def]
    """Editing docx_templates.py without running `build_doc_templates` is caught here."""
    cv, letter = tmp_path / "cv.docx", tmp_path / "letter.docx"
    docx_templates.build_cv_template(cv)
    docx_templates.build_cover_letter_template(letter)
    for built, name in ((cv, "cv_classic.docx"), (letter, "cover_letter.docx")):
        committed = Document(str(documents.TEMPLATE_DIR / name))
        assert [p.text for p in Document(str(built)).paragraphs] == [
            p.text for p in committed.paragraphs
        ]


def test_a_missing_docx_template_is_a_render_error(kit_world, monkeypatch) -> None:  # type: ignore[no-untyped-def]
    from apps.kits.exceptions import RenderError

    kit, _ = kit_world
    monkeypatch.setitem(documents.DOCX_TEMPLATES, documents.DOC_CV, "nope.docx")
    with pytest.raises(RenderError):
        documents.render_docx(documents.DOC_CV, documents.build_cv_context(kit))


# --- PDFs (need WeasyPrint's Pango libraries) ----------------------------------------------------------


def read_pdf(data: bytes) -> Any:
    from pypdf import PdfReader

    return PdfReader(io.BytesIO(data))


@requires_pdf
def test_the_cv_pdf_has_selectable_text_headings_and_database_facts(kit_world) -> None:  # type: ignore[no-untyped-def]
    kit, _ = kit_world
    pdf = documents.html_to_pdf(
        documents.render_html(documents.DOC_CV, documents.build_cv_context(kit))
    )
    assert pdf[:4] == b"%PDF"
    reader = read_pdf(pdf)
    text = "\n".join(page.extract_text() for page in reader.pages)
    for expected in (
        "Jane Wanjiru Doe",
        "SUMMARY",
        "SKILLS",
        "EXPERIENCE",
        "PROJECTS",
        "EDUCATION",
        "CERTIFICATIONS",
        "IT Support Technician, Acme Ltd",
        "Mar 2021 - Present",
        "Resolved support tickets for 120 staff",
        "Nairobi Polytechnic",
    ):
        assert expected in text, expected
    assert 1 <= len(reader.pages) <= 2


@requires_pdf
def test_the_cv_pdf_is_a4_with_embedded_fonts(kit_world) -> None:  # type: ignore[no-untyped-def]
    kit, _ = kit_world
    pdf = documents.html_to_pdf(
        documents.render_html(documents.DOC_CV, documents.build_cv_context(kit))
    )
    page = read_pdf(pdf).pages[0]
    assert round(float(page.mediabox.width)) == 595 and round(float(page.mediabox.height)) == 842
    fonts = page["/Resources"]["/Font"]
    assert fonts, "the PDF uses no fonts"
    for ref in fonts.values():
        font = ref.get_object()
        descriptor = font.get("/FontDescriptor")
        if descriptor is None and "/DescendantFonts" in font:
            descriptor = font["/DescendantFonts"][0].get_object().get("/FontDescriptor")
        assert descriptor is not None
        # embedded: the font program is in the file, so the text copies out correctly
        assert any(key.startswith("/FontFile") for key in descriptor.get_object())
    assert len(list(page.images)) == 0  # no images or icons


@requires_pdf
def test_the_cv_pdf_is_a_single_column(kit_world) -> None:  # type: ignore[no-untyped-def]
    kit, _ = kit_world
    pdf = documents.html_to_pdf(
        documents.render_html(documents.DOC_CV, documents.build_cv_context(kit))
    )
    page = read_pdf(pdf).pages[0]
    starts: list[float] = []

    def visit(text: str, cm: Any, tm: Any, font_dict: Any, font_size: Any) -> None:
        if text.strip():
            starts.append(float(tm[4]))

    page.extract_text(visitor_text=visit)
    width = float(page.mediabox.width)
    assert starts and max(starts) < width * 0.3  # every line begins at the left margin or indent


@requires_pdf
def test_the_cover_letter_pdf_contains_the_letter(kit_world) -> None:  # type: ignore[no-untyped-def]
    kit, _ = kit_world
    pdf = documents.html_to_pdf(
        documents.render_html(documents.DOC_COVER, documents.build_letter_context(kit))
    )
    text = "\n".join(p.extract_text() for p in read_pdf(pdf).pages)
    assert (
        "Dear Hiring Team," in text and "Kind regards," in text and "I keep users working" in text
    )


@requires_pdf
def test_a_long_cv_flows_onto_a_second_page_without_losing_text(kit_world) -> None:  # type: ignore[no-untyped-def]
    kit, ids = kit_world
    cv = good_cv(ids)
    cv["experiences"][0]["bullets"] = [
        {"text": f"Set up laptop batch number {i} for staff.", "source_ids": [ids.bul2]}
        for i in range(5)
    ] * 1
    kit.cv_data = cv
    ctx = documents.build_cv_context(kit)
    ctx["experiences"] = ctx["experiences"] * 4
    pdf = documents.html_to_pdf(documents.render_html(documents.DOC_CV, ctx))
    reader = read_pdf(pdf)
    assert len(reader.pages) >= 2
    assert "CERTIFICATIONS" in "\n".join(p.extract_text() for p in reader.pages)


def test_pdf_rendering_without_pango_is_a_clean_error(monkeypatch) -> None:  # type: ignore[no-untyped-def]
    """Where the libraries are missing the app still imports and reports a user-safe error."""
    import builtins

    from apps.kits.exceptions import RenderError

    real_import = builtins.__import__

    def fake_import(name: str, *args: Any, **kwargs: Any) -> Any:
        if name == "weasyprint":
            raise OSError("cannot load library 'pango'")
        return real_import(name, *args, **kwargs)

    monkeypatch.setattr(builtins, "__import__", fake_import)
    with pytest.raises(RenderError) as excinfo:
        documents.html_to_pdf("<p>x</p>")
    assert "not available" in excinfo.value.message
    assert isinstance(PDF_AVAILABLE, bool)
