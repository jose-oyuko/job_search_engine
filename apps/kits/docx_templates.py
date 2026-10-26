"""Builds the Word templates that docxtpl fills in (spec §12.6).

The `.docx` files are generated here and committed, so the layout is reviewable as code. Tags that
must occupy a paragraph of their own (`{%p ... %}`) are removed with it when rendered. Run
`python manage.py build_doc_templates` after changing this file.
"""

from pathlib import Path

from docx import Document
from docx.document import Document as DocumentType
from docx.shared import Pt

FONT = "Calibri"  # Word users will not have Source Sans 3; Calibri is the safe equivalent


def _base() -> DocumentType:
    document = Document()
    for name in ("Normal", "List Bullet", "Heading 1", "Title"):
        style = document.styles[name]
        style.font.name = FONT
    document.styles["Normal"].font.size = Pt(10.5)
    heading = document.styles["Heading 1"]
    heading.font.size = Pt(11)
    heading.font.bold = True
    return document


def _para(document: DocumentType, text: str, *, bold: bool = False, italic: bool = False,
          size: float | None = None, style: str | None = None) -> None:  # fmt: skip
    paragraph = document.add_paragraph(style=style)
    run = paragraph.add_run(text)
    run.bold, run.italic = bold, italic
    if size:
        run.font.size = Pt(size)


def build_cv_template(path: Path) -> None:
    doc = _base()
    _para(doc, "{{ name }}", bold=True, size=20)
    _para(doc, "{{ headline }}", bold=True, size=11.5)
    _para(doc, "{{ contact_line }}")

    _para(doc, "{%p if summary %}")
    doc.add_heading("Summary", level=1)
    _para(doc, "{{ summary }}")
    _para(doc, "{%p endif %}")

    _para(doc, "{%p if skills %}")
    doc.add_heading("Skills", level=1)
    _para(doc, "{{ skills_line }}")
    _para(doc, "{%p endif %}")

    _para(doc, "{%p if experiences %}")
    doc.add_heading("Experience", level=1)
    _para(doc, "{%p for exp in experiences %}")
    _para(doc, "{{ exp.title }}, {{ exp.company }}", bold=True)
    _para(doc, "{{ exp.dates }}  {{ exp.location }}", italic=True)
    _para(doc, "{%p for bullet in exp.bullets %}")
    _para(doc, "{{ bullet }}", style="List Bullet")
    _para(doc, "{%p endfor %}")
    _para(doc, "{%p endfor %}")
    _para(doc, "{%p endif %}")

    _para(doc, "{%p if projects %}")
    doc.add_heading("Projects", level=1)
    _para(doc, "{%p for project in projects %}")
    _para(doc, "{{ project.name }}  {{ project.role }}", bold=True)
    _para(doc, "{%p for bullet in project.bullets %}")
    _para(doc, "{{ bullet }}", style="List Bullet")
    _para(doc, "{%p endfor %}")
    _para(doc, "{%p endfor %}")
    _para(doc, "{%p endif %}")

    _para(doc, "{%p if education %}")
    doc.add_heading("Education", level=1)
    _para(doc, "{%p for edu in education %}")
    _para(doc, "{{ edu.title }}  {{ edu.institution }}", bold=True)
    _para(doc, "{{ edu.dates }}  {{ edu.grade }}", italic=True)
    _para(doc, "{%p endfor %}")
    _para(doc, "{%p endif %}")

    _para(doc, "{%p if certifications %}")
    doc.add_heading("Certifications", level=1)
    _para(doc, "{%p for cert in certifications %}")
    _para(doc, "{{ cert.name }}  {{ cert.issuer }}  {{ cert.issued }}", style="List Bullet")
    _para(doc, "{%p endfor %}")
    _para(doc, "{%p endif %}")
    doc.save(str(path))


def build_cover_letter_template(path: Path) -> None:
    doc = _base()
    _para(doc, "{{ name }}", bold=True, size=16)
    _para(doc, "{{ contact_line }}")
    _para(doc, "{{ date }}")
    _para(doc, "{{ salutation }}")
    _para(doc, "{%p for paragraph in paragraphs %}")
    _para(doc, "{{ paragraph }}")
    _para(doc, "{%p endfor %}")
    _para(doc, "{{ closing }}")
    _para(doc, "{{ name }}")
    doc.save(str(path))
